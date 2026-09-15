import streamlit as st
import pandas as pd
import datetime
import io
import json
import os
import re
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
try:
    import qrcode
    from PIL import Image
    QR_DISPONIBLE = True
except ImportError:
    QR_DISPONIBLE = False

# =====================================================================
# CONFIGURACIÓN GENERAL Y ESTILOS EN TEMA OSCURO
# =====================================================================
st.set_page_config(
    page_title="Gestión de Préstamos - Laboratorio Universitario",
    page_icon="🔬",
    layout="wide",
    initial_sidebar_state="collapsed"
)

NOMBRE_ARCHIVO_EXCEL = "equipos_laboratorio.xlsx"
ARCHIVO_BD_PRESTAMOS = "historial_prestamos.json"
ARCHIVO_CONFIG_CORREO = "config_correo.json"
ARCHIVO_CONFIG_SISTEMA = "config_sistema.json"

# =====================================================================
# GESTIÓN DE CONFIGURACIÓN DEL SISTEMA (TIEMPO Y CANTIDAD DE EQUIPOS)
# =====================================================================
def cargar_config_sistema():
    defaults = {
        "duracion_prestamo_horas": 2.0,
        "max_equipos_permitidos": 5,
        "nombre_laboratorio": "Laboratorio Universitario - Préstamos"
    }
    if os.path.exists(ARCHIVO_CONFIG_SISTEMA):
        try:
            with open(ARCHIVO_CONFIG_SISTEMA, "r", encoding='utf-8') as f:
                defaults.update(json.load(f))
        except Exception:
            pass
    return defaults

def guardar_config_sistema(config):
    with open(ARCHIVO_CONFIG_SISTEMA, "w", encoding='utf-8') as f:
        json.dump(config, f, indent=4, ensure_ascii=False)

cfg_sistema = cargar_config_sistema()
DURACION_PRESTAMO_HORAS = float(cfg_sistema.get("duracion_prestamo_horas", 2.0))
MAX_EQUIPOS_PERMITIDOS = int(cfg_sistema.get("max_equipos_permitidos", 5))

# =====================================================================
# ESTILOS CSS PROFESIONALES - TEMA OSCURO (DARK THEME)
# =====================================================================
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
    
    html, body, [class*="css"] {
        font-family: 'Inter', sans-serif;
    }
    
    /* Fondo principal oscuro */
    .stApp {
        background-color: #0b0f19;
        color: #f1f5f9;
    }

    /* Encabezado principal moderno con gradiente oscuro */
    .header-banner-dark {
        background: linear-gradient(135deg, #0f172a 0%, #1e3a8a 60%, #1d4ed8 100%);
        border: 1px solid #1e3a8a;
        color: white;
        padding: 24px 30px;
        border-radius: 12px;
        margin-bottom: 24px;
        box-shadow: 0 4px 20px rgba(0, 0, 0, 0.4);
    }
    .header-banner-dark h1 {
        color: #ffffff !important;
        font-size: 26px;
        font-weight: 700;
        margin: 0;
        display: flex;
        align-items: center;
        gap: 12px;
    }
    .header-banner-dark p {
        color: #93c5fd;
        font-size: 14px;
        margin: 6px 0 0 0;
    }

    /* Tarjetas de métricas oscuras */
    .metric-card-dark {
        background: #151e2e;
        border: 1px solid #1e293b;
        border-radius: 10px;
        padding: 18px 14px;
        text-align: center;
        box-shadow: 0 4px 12px rgba(0,0,0,0.25);
        transition: transform 0.2s ease, border-color 0.2s ease;
    }
    .metric-card-dark:hover {
        transform: translateY(-2px);
        border-color: #3b82f6;
    }
    .metric-value-dark {
        font-size: 28px;
        font-weight: 700;
        color: #f8fafc;
        margin-bottom: 4px;
    }
    .metric-label-dark {
        font-size: 11px;
        font-weight: 600;
        text-transform: uppercase;
        letter-spacing: 0.6px;
        color: #94a3b8;
    }
    .metric-alert-dark {
        border: 1px solid #ef4444 !important;
        border-left: 5px solid #ef4444 !important;
        background: rgba(239, 68, 68, 0.1) !important;
    }
    .metric-alert-dark .metric-value-dark {
        color: #f87171 !important;
    }

    /* Badges de estado en tema oscuro */
    .badge {
        display: inline-block;
        padding: 4px 12px;
        border-radius: 9999px;
        font-size: 12px;
        font-weight: 600;
    }
    .badge-vencido {
        background-color: rgba(239, 68, 68, 0.2);
        color: #fca5a5;
        border: 1px solid #ef4444;
        animation: pulse 2s cubic-bezier(0.4, 0, 0.6, 1) infinite;
    }
    .badge-activo {
        background-color: rgba(59, 130, 246, 0.2);
        color: #93c5fd;
        border: 1px solid #3b82f6;
    }
    .badge-pendiente {
        background-color: rgba(245, 158, 11, 0.2);
        color: #fde047;
        border: 1px solid #eab308;
    }
    .badge-finalizado {
        background-color: rgba(16, 185, 129, 0.2);
        color: #86efac;
        border: 1px solid #10b981;
    }
    .badge-cancelado {
        background-color: rgba(100, 116, 139, 0.2);
        color: #cbd5e1;
        border: 1px solid #475569;
    }

    @keyframes pulse {
        0%, 100% { opacity: 1; }
        50% { opacity: 0.6; }
    }

    /* Caja de alerta para préstamos vencidos */
    .overdue-alert-dark {
        background: rgba(239, 68, 68, 0.15);
        border: 1px solid #ef4444;
        border-left: 6px solid #ef4444;
        border-radius: 8px;
        padding: 14px 18px;
        margin: 12px 0;
        color: #fca5a5;
        font-weight: 500;
    }

    /* Tarjetas informativas de fondo oscuro */
    .info-box-dark {
        background: #151e2e;
        border: 1px solid #1e293b;
        border-radius: 10px;
        padding: 20px;
        color: #cbd5e1;
        box-shadow: 0 4px 12px rgba(0,0,0,0.2);
    }
    .info-box-dark h4 {
        color: #60a5fa !important;
        margin-top: 0;
    }

    /* Pestañas en tema oscuro */
    .stTabs [data-baseweb="tab-list"] {
        gap: 8px;
        border-bottom: 2px solid #1e293b;
        margin-bottom: 20px;
    }
    .stTabs [data-baseweb="tab"] {
        padding: 10px 20px;
        border-radius: 8px 8px 0 0;
        font-weight: 600;
        color: #94a3b8;
    }
    .stTabs [aria-selected="true"] {
        background-color: #1e293b !important;
        color: #38bdf8 !important;
        border-bottom: 3px solid #38bdf8 !important;
    }
</style>
""", unsafe_allow_html=True)

# =====================================================================
# 1. GESTIÓN Y NORMALIZACIÓN DE INVENTARIO (EXCEL)
# =====================================================================
def crear_inventario_por_defecto():
    df = pd.DataFrame({
        "codigo_equipo": [
            "LAB-OSC-01", "LAB-OSC-02", "LAB-MUL-01", "LAB-MUL-02",
            "LAB-FUE-01", "LAB-GEN-01", "LAB-CAU-01", "LAB-FPG-01",
            "LAB-ARD-01", "LAB-MIC-01"
        ],
        "nombre_equipo": [
            "Osciloscopio Digital 100MHz", "Osciloscopio Digital 50MHz",
            "Multímetro Digital True RMS", "Multímetro de Banco",
            "Fuente Regulada DC 30V/5A", "Generador de Señales Arbitrarias",
            "Estación de Soldadura Cautín", "Tarjeta de Desarrollo FPGA Cyclone IV",
            "Kit Arduino Uno + Sensores Básicos", "Microscopio Digital para PCB"
        ],
        "categoria": [
            "Medición", "Medición", "Medición", "Medición",
            "Alimentación", "Generación", "Herramientas", "Embebidos",
            "Embebidos", "Óptica"
        ],
        "estado": ["Disponible"] * 10,
        "ubicacion": [
            "Estante A1", "Estante A2", "Estante B1", "Estante B2",
            "Mesa 1", "Mesa 2", "Mesa 3", "Cajón C1", "Cajón C2", "Mesa 4"
        ]
    })
    df.to_excel(NOMBRE_ARCHIVO_EXCEL, index=False)
    return df

@st.cache_data(ttl=10)
def cargar_equipos():
    if not os.path.exists(NOMBRE_ARCHIVO_EXCEL):
        return crear_inventario_por_defecto()

    try:
        wb = openpyxl.load_workbook(NOMBRE_ARCHIVO_EXCEL, read_only=True)
        sheet_names = wb.sheetnames
        wb.close()

        if "ArmarioLCD" in sheet_names:
            df_raw = pd.read_excel(NOMBRE_ARCHIVO_EXCEL, sheet_name="ArmarioLCD", header=6)
            
            col_desc = None
            for col in df_raw.columns:
                c_str = str(col).lower()
                if "art" in c_str or "descrip" in c_str or "nombre" in c_str:
                    col_desc = col
                    break
            
            if col_desc is None and len(df_raw.columns) >= 2:
                col_desc = df_raw.columns[1]

            col_cat = None
            for col in df_raw.columns:
                c_str = str(col).lower()
                if "armario" in c_str or "cat" in c_str or "ubic" in c_str:
                    col_cat = col
                    break
            if col_cat is None and len(df_raw.columns) >= 4:
                col_cat = df_raw.columns[3]

            df_raw = df_raw.dropna(subset=[col_desc])
            df_raw = df_raw[df_raw[col_desc].astype(str).str.strip() != ""]

            registros = []
            for idx, row in df_raw.iterrows():
                nombre = str(row[col_desc]).strip()
                cat = str(row[col_cat]).strip() if (col_cat and pd.notna(row[col_cat])) else "General"
                registros.append({
                    "codigo_equipo": f"LCD-{len(registros)+1:03d}",
                    "nombre_equipo": nombre,
                    "categoria": cat,
                    "estado": "Disponible",
                    "ubicacion": cat
                })
            
            df_normalizado = pd.DataFrame(registros)
            if not df_normalizado.empty:
                return df_normalizado

        df_std = pd.read_excel(NOMBRE_ARCHIVO_EXCEL, sheet_name=0)
        df_std.columns = [str(c).strip().lower() for c in df_std.columns]

        map_cols = {}
        for c in df_std.columns:
            if c in ["codigo_equipo", "codigo", "id", "código", "cod"]:
                map_cols[c] = "codigo_equipo"
            elif c in ["nombre_equipo", "nombre", "equipo", "articulo", "artículo", "descripcion", "descripción"]:
                map_cols[c] = "nombre_equipo"
            elif c in ["categoria", "categoría", "tipo", "armario"]:
                map_cols[c] = "categoria"
            elif c in ["estado", "status", "disponibilidad"]:
                map_cols[c] = "estado"
            elif c in ["ubicacion", "ubicación", "estante", "mesa"]:
                map_cols[c] = "ubicacion"

        df_std = df_std.rename(columns=map_cols)

        if "nombre_equipo" not in df_std.columns:
            if len(df_std.columns) > 0:
                df_std["nombre_equipo"] = df_std.iloc[:, 0].astype(str)
            else:
                return crear_inventario_por_defecto()

        if "codigo_equipo" not in df_std.columns:
            df_std["codigo_equipo"] = [f"EQ-{i+1:03d}" for i in range(len(df_std))]
        if "categoria" not in df_std.columns:
            df_std["categoria"] = "General"
        if "estado" not in df_std.columns:
            df_std["estado"] = "Disponible"
        if "ubicacion" not in df_std.columns:
            df_std["ubicacion"] = "Laboratorio"

        df_final = df_std[["codigo_equipo", "nombre_equipo", "categoria", "estado", "ubicacion"]].dropna(subset=["nombre_equipo"])
        if df_final.empty:
            return crear_inventario_por_defecto()
        return df_final

    except Exception as err:
        st.warning(f"⚠️ No se pudo leer el archivo Excel ({err}). Generando plantilla de respaldo...")
        return crear_inventario_por_defecto()

# =====================================================================
# 2. PERSISTENCIA DE PRÉSTAMOS (JSON)
# =====================================================================
def cargar_prestamos():
    if os.path.exists(ARCHIVO_BD_PRESTAMOS):
        try:
            with open(ARCHIVO_BD_PRESTAMOS, "r", encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            return []
    return []

def guardar_prestamos(datos):
    with open(ARCHIVO_BD_PRESTAMOS, "w", encoding='utf-8') as f:
        json.dump(datos, f, indent=4, ensure_ascii=False)

if "prestamos_db" not in st.session_state:
    st.session_state.prestamos_db = cargar_prestamos()

def sincronizar_prestamos():
    st.session_state.prestamos_db = cargar_prestamos()

# =====================================================================
# 3. GESTIÓN Y ENVÍO DE CORREOS CON GMAIL (SMTP)
# =====================================================================
def cargar_config_correo():
    defaults = {
        "proveedor": "gmail",
        "smtp_server": "smtp.gmail.com",
        "smtp_port": 465,
        "usar_ssl": True,
        "remitente": "",
        "password": "",
        "nombre_remitente": "Laboratorio Universitario - Préstamos"
    }
    if os.path.exists(ARCHIVO_CONFIG_CORREO):
        try:
            with open(ARCHIVO_CONFIG_CORREO, "r", encoding='utf-8') as f:
                defaults.update(json.load(f))
        except Exception:
            pass
    return defaults

def guardar_config_correo(config):
    with open(ARCHIVO_CONFIG_CORREO, "w", encoding='utf-8') as f:
        json.dump(config, f, indent=4, ensure_ascii=False)

def enviar_correo_confirmacion(destinatario, nombre_estudiante, cedula, equipos, id_solicitud, duracion_horas=2.0):
    config = cargar_config_correo()
    remitente = config.get("remitente", "").strip()
    password = config.get("password", "").strip().replace(" ", "")
    smtp_server = config.get("smtp_server", "smtp.gmail.com").strip()
    smtp_port = int(config.get("smtp_port", 465))
    usar_ssl = config.get("usar_ssl", True)
    nombre_remitente = config.get("nombre_remitente", "Laboratorio Universitario - Préstamos")

    if not remitente or not password:
        return False, "Faltan credenciales de Gmail (correo remitente o contraseña de aplicación de 16 letras)."

    msg = MIMEMultipart('alternative')
    msg['From'] = f"{nombre_remitente} <{remitente}>"
    msg['To'] = destinatario
    msg['Subject'] = f"✅ Confirmación de Préstamo #{id_solicitud} - Laboratorio Universitario"

    fecha_hoy = datetime.datetime.now().strftime("%d/%m/%Y %I:%M %p")
    lista_equipos_txt = "\n- ".join(equipos)

    cuerpo_texto = f"""Hola {nombre_estudiante},

Tu solicitud de préstamo de equipos #{id_solicitud} ha sido registrada exitosamente en el Laboratorio.

Detalles:
- Cédula / Código: {cedula}
- Fecha y hora: {fecha_hoy}
- Material solicitado ({len(equipos)} artículos):
- {lista_equipos_txt}

Por favor, acércate al mostrador del laboratorio con tu documento de identidad para retirar los equipos. El tiempo estándar de uso es de {duracion_horas} horas.

Atentamente,
{nombre_remitente}
"""

    items_html = "".join([
        f"""<li style="margin-bottom: 8px; color: #1e293b; font-size: 14px;">
            <span style="color: #2563eb; font-weight: bold; margin-right: 6px;">✔</span>
            <strong>{eq}</strong>
        </li>""" for eq in equipos
    ])

    cuerpo_html = f"""
    <!DOCTYPE html>
    <html lang="es">
    <head><meta charset="utf-8"></head>
    <body style="font-family: 'Segoe UI', Arial, sans-serif; background-color: #f1f5f9; margin: 0; padding: 24px;">
        <div style="max-width: 580px; margin: 0 auto; background-color: #ffffff; border-radius: 12px; border: 1px solid #e2e8f0; overflow: hidden; box-shadow: 0 4px 12px rgba(0,0,0,0.05);">
            <div style="background: linear-gradient(135deg, #1e3a8a 0%, #2563eb 100%); padding: 28px; text-align: center; color: #ffffff;">
                <h1 style="margin: 0; font-size: 22px; font-weight: 700; letter-spacing: -0.5px;">🔬 Laboratorio Universitario</h1>
                <p style="margin: 6px 0 0 0; font-size: 14px; opacity: 0.9;">Comprobante Oficial de Préstamo de Equipos</p>
            </div>
            <div style="padding: 28px;">
                <p style="font-size: 16px; color: #0f172a; margin-top: 0;">Hola <strong>{nombre_estudiante}</strong>,</p>
                <p style="color: #475569; font-size: 14px; line-height: 1.6;">
                    Tu solicitud <strong>#{id_solicitud}</strong> ha sido registrada en el sistema. A continuación el detalle de los materiales solicitados:
                </p>
                
                <table style="width: 100%; margin: 16px 0; border-collapse: collapse; font-size: 13px;">
                    <tr style="border-bottom: 1px solid #e2e8f0;">
                        <td style="padding: 8px 0; color: #64748b; font-weight: 600;">Cédula / Código:</td>
                        <td style="padding: 8px 0; color: #0f172a; font-weight: bold; text-align: right;">{cedula}</td>
                    </tr>
                    <tr style="border-bottom: 1px solid #e2e8f0;">
                        <td style="padding: 8px 0; color: #64748b; font-weight: 600;">Fecha de Solicitud:</td>
                        <td style="padding: 8px 0; color: #0f172a; text-align: right;">{fecha_hoy}</td>
                    </tr>
                    <tr>
                        <td style="padding: 8px 0; color: #64748b; font-weight: 600;">Tiempo Asignado de Préstamo:</td>
                        <td style="padding: 8px 0; color: #0f172a; text-align: right;"><strong>{duracion_horas} Horas</strong></td>
                    </tr>
                </table>

                <div style="background-color: #f8fafc; border: 1px solid #cbd5e1; border-left: 4px solid #2563eb; padding: 16px; border-radius: 8px; margin: 20px 0;">
                    <div style="font-size: 13px; font-weight: 700; color: #1e3a8a; text-transform: uppercase; margin-bottom: 8px; letter-spacing: 0.5px;">Equipos Solicitados ({len(equipos)})</div>
                    <ul style="margin: 0; padding-left: 4px; list-style-type: none;">
                        {items_html}
                    </ul>
                </div>

                <div style="background-color: #fefce8; border: 1px solid #fef08a; border-radius: 8px; padding: 12px 16px; margin-top: 20px; font-size: 13px; color: #854d0e;">
                    📍 <strong>Paso siguiente:</strong> Acércate al mostrador de entrega con tu carnet institucional para retirar tu material.
                </div>
            </div>
            <div style="background-color: #f8fafc; border-top: 1px solid #e2e8f0; padding: 16px; text-align: center; font-size: 12px; color: #94a3b8;">
                {nombre_remitente} · Notificación generada automáticamente
            </div>
        </div>
    </body>
    </html>
    """

    msg.attach(MIMEText(cuerpo_texto, 'plain', 'utf-8'))
    msg.attach(MIMEText(cuerpo_html, 'html', 'utf-8'))

    try:
        if usar_ssl or smtp_port == 465:
            server = smtplib.SMTP_SSL(smtp_server, smtp_port, timeout=12)
        else:
            server = smtplib.SMTP(smtp_server, smtp_port, timeout=12)
            server.ehlo()
            server.starttls()
            server.ehlo()

        server.login(remitente, password)
        server.send_message(msg)
        server.quit()
        return True, ""
    except Exception as e:
        err = str(e)
        if "535" in err or "authentication" in err.lower() or "badcredentials" in err.lower():
            return False, "Error de autenticación en Gmail (535): Asegúrate de usar una Contraseña de Aplicación de 16 letras generada en Google y tener la verificación en 2 pasos activa."
        elif "timed out" in err.lower() or "timeout" in err.lower():
            return False, f"Tiempo de espera agotado al conectar con {smtp_server}:{smtp_port}."
        elif "550" in err or "limit" in err.lower():
            return False, "Límite diario de envíos de Gmail alcanzado."
        return False, err

# =====================================================================
# 4. EXPORTACIÓN PROFESIONAL A EXCEL
# =====================================================================
def exportar_historial_excel(prestamos):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historial de Préstamos"

    headers = [
        "ID", "Cédula", "Estudiante", "Correo", "Equipos Solicitados",
        "Estado", "Hora Solicitud", "Hora Entrega", "Hora Límite",
        "Devolución Real", "Retraso (Minutos)"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for p in prestamos:
        retraso = 0
        if p.get("estado") == "Activo (Entregado)" and "timestamp_limite_iso" in p:
            ahora = datetime.datetime.now()
            limite = datetime.datetime.fromisoformat(p["timestamp_limite_iso"])
            if ahora > limite:
                retraso = int((ahora - limite).total_seconds() / 60)
        elif p.get("minutos_retraso_final"):
            retraso = p.get("minutos_retraso_final")

        equipos_str = ", ".join(p.get("equipos", []))
        row = [
            p.get("id"),
            p.get("cedula"),
            p.get("nombre"),
            p.get("correo"),
            equipos_str,
            p.get("estado"),
            p.get("hora_solicitud", "N/A"),
            p.get("hora_entrega_texto", "N/A"),
            p.get("hora_limite_texto", "N/A"),
            p.get("hora_devolucion_real", "N/A"),
            retraso if retraso > 0 else 0
        ]
        ws.append(row)

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            cell.border = thin_border
            if cell.row > 1:
                cell.font = Font(name="Segoe UI", size=10)
                cell.alignment = Alignment(vertical="center")
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

    ws.row_dimensions[1].height = 28

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

# =====================================================================
# 5. CARGA DE DATOS PRINCIPALES
# =====================================================================
df_equipos = cargar_equipos()
sincronizar_prestamos()

# Lista de todas las etiquetas de equipos en el laboratorio
opciones_totales_lab = []
for _, row in df_equipos.iterrows():
    codigo = row.get("codigo_equipo", "")
    nombre = row.get("nombre_equipo", "")
    eq_label = f"{codigo} — {nombre}" if codigo else nombre
    opciones_totales_lab.append(eq_label)

equipos_ocupados = set()
for p in st.session_state.prestamos_db:
    if p.get("estado") in ["Pendiente", "Activo (Entregado)"]:
        for eq in p.get("equipos", []):
            equipos_ocupados.add(eq)

# =====================================================================
# ENCABEZADO PRINCIPAL DE LA APLICACIÓN (DARK THEME)
# =====================================================================
st.markdown(f"""
<div class="header-banner-dark">
    <h1>🔬 Sistema de Préstamos de Laboratorio</h1>
    <p>Gestión y control en tiempo real · Tiempo configurado: <strong>{DURACION_PRESTAMO_HORAS} horas</strong> · Máximo por solicitud: <strong>{MAX_EQUIPOS_PERMITIDOS} equipos</strong> · Modo Oscuro</p>
</div>
""", unsafe_allow_html=True)

tab_estudiante, tab_encargado, tab_inventario, tab_config = st.tabs([
    "📱 Solicitud de Estudiantes",
    "🖥️ Control del Encargado",
    "📊 Inventario de Equipos",
    "⚙️ Menú de Configuración (Tiempo & Cantidades)"
])

# =====================================================================
# PESTAÑA 1: VISTA DEL ESTUDIANTE (SOLICITUD)
# =====================================================================
with tab_estudiante:
    st.markdown("### 📝 Formulario de Solicitud de Préstamo")
    st.markdown(f"Completa tus datos personales y selecciona los equipos que requieres (puedes solicitar hasta **{MAX_EQUIPOS_PERMITIDOS} equipos**).")

    col_form, col_info = st.columns([3, 2], gap="large")

    with col_form:
        with st.form("form_prestamo_estudiante", clear_on_submit=False):
            st.markdown("##### 👤 Datos del Estudiante")
            c1, c2 = st.columns(2)
            with c1:
                cedula_input = st.text_input("Cédula / Código Institucional *", placeholder="Ej: 1098765432")
            with c2:
                nombre_input = st.text_input("Nombre Completo *", placeholder="Ej: Andrea Gómez Torres")
            
            correo_input = st.text_input("Correo Institucional / Personal (Gmail) *", placeholder="Ej: estudiante@gmail.com o usuario@universidad.edu.co")

            st.markdown("##### 🔍 Selección de Equipos de Laboratorio")
            
            categorias = ["Todas las categorías"] + sorted(df_equipos["categoria"].dropna().unique().tolist())
            categoria_seleccionada = st.selectbox("Filtrar por categoría / armario:", options=categorias)

            if categoria_seleccionada != "Todas las categorías":
                df_filtrado = df_equipos[df_equipos["categoria"] == categoria_seleccionada]
            else:
                df_filtrado = df_equipos

            opciones_disponibles = []
            for _, row in df_filtrado.iterrows():
                codigo = row.get("codigo_equipo", "")
                nombre = row.get("nombre_equipo", "")
                eq_label = f"{codigo} — {nombre}" if codigo else nombre
                if eq_label not in equipos_ocupados:
                    opciones_disponibles.append(eq_label)

            equipos_seleccionados = st.multiselect(
                f"Equipos a solicitar (Límite configurado: hasta {MAX_EQUIPOS_PERMITIDOS} equipos):",
                options=opciones_disponibles,
                max_selections=MAX_EQUIPOS_PERMITIDOS,
                help=f"Solo se muestran equipos actualmente disponibles en bodega. Puedes elegir hasta {MAX_EQUIPOS_PERMITIDOS} elementos."
            )

            observaciones_input = st.text_area("Observaciones o motivo de la práctica (opcional):", placeholder="Ej: Práctica de Circuitos Digitales II - Mesa 3")

            btn_solicitar = st.form_submit_button("🚀 Enviar Solicitud de Préstamo", use_container_width=True)

            if btn_solicitar:
                if not cedula_input.strip() or not nombre_input.strip() or not correo_input.strip():
                    st.error("❌ Por favor completa todos los campos obligatorios (*) del estudiante.")
                elif "@" not in correo_input or "." not in correo_input:
                    st.error("❌ Por favor ingresa una dirección de correo electrónico válida.")
                elif not equipos_seleccionados:
                    st.error("❌ Debes seleccionar al menos un equipo disponible de la lista.")
                elif len(equipos_seleccionados) > MAX_EQUIPOS_PERMITIDOS:
                    st.error(f"❌ Has superado el límite de {MAX_EQUIPOS_PERMITIDOS} equipos permitidos.")
                else:
                    nuevo_id = len(st.session_state.prestamos_db) + 1
                    ahora = datetime.datetime.now()
                    nueva_solicitud = {
                        "id": nuevo_id,
                        "cedula": cedula_input.strip(),
                        "nombre": nombre_input.strip(),
                        "correo": correo_input.strip(),
                        "equipos": equipos_seleccionados,
                        "observaciones": observaciones_input.strip(),
                        "hora_solicitud": ahora.strftime("%d/%m/%Y %I:%M %p"),
                        "timestamp_solicitud": ahora.isoformat(),
                        "duracion_asignada_horas": DURACION_PRESTAMO_HORAS,
                        "estado": "Pendiente"
                    }

                    st.session_state.prestamos_db.append(nueva_solicitud)
                    guardar_prestamos(st.session_state.prestamos_db)

                    with st.spinner("Enviando comprobante oficial desde Gmail..."):
                        correo_ok, err_msg = enviar_correo_confirmacion(
                            destinatario=correo_input.strip(),
                            nombre_estudiante=nombre_input.strip(),
                            cedula=cedula_input.strip(),
                            equipos=equipos_seleccionados,
                            id_solicitud=nuevo_id,
                            duracion_horas=DURACION_PRESTAMO_HORAS
                        )

                    st.success(f"🎉 ¡Solicitud #{nuevo_id} registrada con éxito!")
                    if correo_ok:
                        st.info(f"📧 Se envió un comprobante oficial de confirmación a **{correo_input}**.")
                    else:
                        st.warning(f"⚠️ La solicitud se guardó en la base de datos, pero el correo falló ({err_msg}). Revisa la pestaña de Configuración.")
                    
                    st.balloons()
                    st.rerun()

    with col_info:
        st.markdown(f"""
        <div class="info-box-dark">
            <h4>📌 Parámetros y Normas Actuales</h4>
            <ul style="font-size: 13px; line-height: 1.8; padding-left: 18px; margin-bottom: 0;">
                <li>Presentar carnet o documento de identidad en el mostrador.</li>
                <li>Tiempo de préstamo estándar configurado: <strong style="color: #38bdf8;">{DURACION_PRESTAMO_HORAS} horas</strong>.</li>
                <li>Límite de equipos a prestar por solicitud: <strong style="color: #38bdf8;">{MAX_EQUIPOS_PERMITIDOS} artículos</strong>.</li>
                <li>Verifica que el material esté en perfecto estado antes de retirarte.</li>
                <li>En caso de daño o extravío, notificar de inmediato al encargado.</li>
                <li>Recibirás un comprobante en tu correo electrónico de Gmail.</li>
            </ul>
        </div>
        """, unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)
        st.markdown("##### 📦 Disponibilidad Rápida")
        total_items = len(df_equipos)
        prestados_count = len(equipos_ocupados)
        disponibles_count = max(total_items - prestados_count, 0)
        
        m_c1, m_c2 = st.columns(2)
        m_c1.metric("Equipos Disponibles", disponibles_count)
        m_c2.metric("En Préstamo", prestados_count)

# =====================================================================
# PESTAÑA 2: VISTA DEL ENCARGADO (CONTROL DE PRÉSTAMOS)
# =====================================================================
with tab_encargado:
    st.markdown("### 🖥️ Tablero de Control de Préstamos en Tiempo Real")

    ahora_dt = datetime.datetime.now()
    prestamos_lista = st.session_state.prestamos_db

    conteo_pendientes = 0
    conteo_activos = 0
    conteo_vencidos = 0
    conteo_finalizados_hoy = 0

    for p in prestamos_lista:
        est = p.get("estado")
        if est == "Pendiente":
            conteo_pendientes += 1
        elif est == "Activo (Entregado)":
            conteo_activos += 1
            if "timestamp_limite_iso" in p:
                limite_obj = datetime.datetime.fromisoformat(p["timestamp_limite_iso"])
                if ahora_dt > limite_obj:
                    conteo_vencidos += 1
        elif est == "Finalizado":
            conteo_finalizados_hoy += 1

    kpi1, kpi2, kpi3, kpi4, kpi5 = st.columns(5)
    with kpi1:
        st.markdown(f"""
        <div class="metric-card-dark">
            <div class="metric-value-dark">{len(prestamos_lista)}</div>
            <div class="metric-label-dark">Total Préstamos</div>
        </div>
        """, unsafe_allow_html=True)
    with kpi2:
        st.markdown(f"""
        <div class="metric-card-dark">
            <div class="metric-value-dark" style="color: #fde047;">{conteo_pendientes}</div>
            <div class="metric-label-dark">🟡 Pendientes</div>
        </div>
        """, unsafe_allow_html=True)
    with kpi3:
        st.markdown(f"""
        <div class="metric-card-dark">
            <div class="metric-value-dark" style="color: #60a5fa;">{conteo_activos}</div>
            <div class="metric-label-dark">🔵 En Uso (Activos)</div>
        </div>
        """, unsafe_allow_html=True)
    with kpi4:
        card_class = "metric-card-dark metric-alert-dark" if conteo_vencidos > 0 else "metric-card-dark"
        st.markdown(f"""
        <div class="{card_class}">
            <div class="metric-value-dark">{conteo_vencidos}</div>
            <div class="metric-label-dark">🚨 Retrasados / Vencidos</div>
        </div>
        """, unsafe_allow_html=True)
    with kpi5:
        st.markdown(f"""
        <div class="metric-card-dark">
            <div class="metric-value-dark" style="color: #4ade80;">{conteo_finalizados_hoy}</div>
            <div class="metric-label-dark">🟢 Finalizados</div>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)

    f_c1, f_c2, f_c3 = st.columns([2, 2, 2])
    with f_c1:
        filtro_estado = st.selectbox(
            "Filtrar por estado:",
            ["Todos", "🚨 Vencidos", "🟡 Pendientes", "🔵 En Uso (Activos)", "🟢 Finalizados", "⚪ Cancelados"]
        )
    with f_c2:
        busqueda_texto = st.text_input("🔍 Buscar por Cédula o Nombre:", placeholder="Ej: 10987... o Carlos")
    with f_c3:
        st.write("")
        st.write("")
        if st.button("🔄 Actualizar Tablero"):
            sincronizar_prestamos()
            st.rerun()

    if not prestamos_lista:
        st.info("ℹ️ No hay registros de préstamos todavía.")
    else:
        registros_filtrados = []
        for idx, p in enumerate(reversed(prestamos_lista)):
            idx_real = len(prestamos_lista) - 1 - idx
            estado = p.get("estado", "")

            es_vencido = False
            minutos_retraso = 0
            if estado == "Activo (Entregado)" and "timestamp_limite_iso" in p:
                limite_obj = datetime.datetime.fromisoformat(p["timestamp_limite_iso"])
                if ahora_dt > limite_obj:
                    es_vencido = True
                    minutos_retraso = int((ahora_dt - limite_obj).total_seconds() / 60)

            if filtro_estado == "🚨 Vencidos" and not es_vencido:
                continue
            elif filtro_estado == "🟡 Pendientes" and estado != "Pendiente":
                continue
            elif filtro_estado == "🔵 En Uso (Activos)" and estado != "Activo (Entregado)":
                continue
            elif filtro_estado == "🟢 Finalizados" and estado != "Finalizado":
                continue
            elif filtro_estado == "⚪ Cancelados" and estado != "Cancelado":
                continue

            if busqueda_texto.strip():
                query = busqueda_texto.strip().lower()
                ced = str(p.get("cedula", "")).lower()
                nom = str(p.get("nombre", "")).lower()
                if query not in ced and query not in nom:
                    continue

            registros_filtrados.append((idx_real, p, es_vencido, minutos_retraso))

        st.markdown(f"**Mostrando {len(registros_filtrados)} registros:**")

        for idx_real, p, es_vencido, minutos_retraso in registros_filtrados:
            estado = p.get("estado", "")
            
            if es_vencido:
                titulo = f"🚨 VENCIDO (+{minutos_retraso}m) — #{p['id']} · {p['nombre']} (Cédula: {p['cedula']})"
            elif estado == "Activo (Entregado)":
                titulo = f"🔵 EN USO — #{p['id']} · {p['nombre']} (Cédula: {p['cedula']})"
            elif estado == "Pendiente":
                titulo = f"🟡 PENDIENTE — #{p['id']} · {p['nombre']} (Cédula: {p['cedula']})"
            elif estado == "Finalizado":
                titulo = f"🟢 FINALIZADO — #{p['id']} · {p['nombre']}"
            else:
                titulo = f"⚪ CANCELADO — #{p['id']} · {p['nombre']}"

            with st.expander(titulo, expanded=(es_vencido or estado == "Pendiente")):
                if es_vencido:
                    st.markdown(f"""
                    <div class="overdue-alert-dark">
                        ⚠️ <strong>ALERTA DE TIEMPO VENCIDO:</strong> Este préstamo superó el límite por <strong>{minutos_retraso} minutos</strong>.<br>
                        Hora Límite Programada: <strong>{p.get('hora_limite_texto', 'N/A')}</strong>.
                    </div>
                    """, unsafe_allow_html=True)

                c_info1, c_info2 = st.columns([3, 2])
                with c_info1:
                    st.write(f"**Estudiante:** {p.get('nombre')} | **Cédula:** {p.get('cedula')}")
                    st.write(f"**Correo:** `{p.get('correo')}`")
                    st.write(f"**Equipos Asignados ({len(p.get('equipos', []))}):** {', '.join(p.get('equipos', []))}")
                    if p.get("observaciones"):
                        st.write(f"**Observaciones:** _{p.get('observaciones')}_")

                with c_info2:
                    st.write(f"**Solicitud:** {p.get('hora_solicitud', 'N/A')}")
                    if "hora_entrega_texto" in p:
                        st.write(f"**Entrega:** {p.get('hora_entrega_texto')}")
                        st.write(f"**Hora Límite:** {p.get('hora_limite_texto')}")
                    if "hora_devolucion_real" in p:
                        st.write(f"**Devuelto:** {p.get('hora_devolucion_real')}")

                st.markdown("---")
                
                if estado == "Pendiente":
                    st.markdown("##### 🛠️ Modificar Parámetros y Equipos antes de la Entrega")
                    
                    # Permite al encargado ajustar la lista y cantidad de equipos a prestar
                    equipos_actuales = [eq for eq in p.get("equipos", []) if eq in opciones_totales_lab]
                    equipos_a_entregar_editados = st.multiselect(
                        "Equipos a entregar (puedes añadir o retirar equipos antes de confirmar):",
                        options=opciones_totales_lab,
                        default=equipos_actuales,
                        key=f"edit_equipos_entrega_{idx_real}",
                        help="Modifica aquí los artículos si necesitas agregar o cambiar alguno."
                    )

                    col_tiempo_ent, col_btn_ent, col_btn_canc = st.columns([2, 2, 2])
                    
                    with col_tiempo_ent:
                        horas_entrega_custom = st.number_input(
                            "Tiempo a otorgar (Horas):",
                            value=float(p.get("duracion_asignada_horas", DURACION_PRESTAMO_HORAS)),
                            min_value=0.25,
                            max_value=24.0,
                            step=0.5,
                            key=f"custom_horas_{idx_real}",
                            help="Modifica el tiempo de préstamo específico para este estudiante."
                        )

                    with col_btn_ent:
                        st.write("")
                        st.write("")
                        if st.button("📦 Entregar Material", key=f"btn_entregar_{idx_real}", use_container_width=True):
                            if not equipos_a_entregar_editados:
                                st.error("Debes incluir al menos 1 equipo para entregar.")
                            else:
                                hora_entrega = datetime.datetime.now()
                                hora_limite = hora_entrega + datetime.timedelta(hours=horas_entrega_custom)
                                
                                st.session_state.prestamos_db[idx_real]["estado"] = "Activo (Entregado)"
                                st.session_state.prestamos_db[idx_real]["equipos"] = equipos_a_entregar_editados
                                st.session_state.prestamos_db[idx_real]["duracion_asignada_horas"] = horas_entrega_custom
                                st.session_state.prestamos_db[idx_real]["timestamp_limite_iso"] = hora_limite.isoformat()
                                st.session_state.prestamos_db[idx_real]["hora_entrega_texto"] = hora_entrega.strftime("%d/%m/%Y %I:%M %p")
                                st.session_state.prestamos_db[idx_real]["hora_limite_texto"] = hora_limite.strftime("%I:%M %p")
                                
                                guardar_prestamos(st.session_state.prestamos_db)
                                st.success(f"Material entregado ({len(equipos_a_entregar_editados)} equipos por {horas_entrega_custom}h).")
                                st.rerun()

                    with col_btn_canc:
                        st.write("")
                        st.write("")
                        if st.button("❌ Cancelar Solicitud", key=f"btn_cancelar_{idx_real}", use_container_width=True):
                            st.session_state.prestamos_db[idx_real]["estado"] = "Cancelado"
                            guardar_prestamos(st.session_state.prestamos_db)
                            st.info(f"Solicitud #{p['id']} cancelada.")
                            st.rerun()

                elif estado == "Activo (Entregado)":
                    btn_c1, _ = st.columns([2, 4])
                    with btn_c1:
                        if st.button("🔄 Registrar Devolución", key=f"btn_devolver_{idx_real}", use_container_width=True):
                            hora_dev = datetime.datetime.now()
                            st.session_state.prestamos_db[idx_real]["estado"] = "Finalizado"
                            st.session_state.prestamos_db[idx_real]["hora_devolucion_real"] = hora_dev.strftime("%d/%m/%Y %I:%M %p")
                            
                            if es_vencido:
                                st.session_state.prestamos_db[idx_real]["minutos_retraso_final"] = minutos_retraso
                            
                            guardar_prestamos(st.session_state.prestamos_db)
                            st.success(f"✅ Devolución registrada para préstamo #{p['id']}.")
                            st.rerun()

    st.markdown("---")
    st.markdown("### 📊 Exportar Historial Completo a Excel")
    st.markdown("Descarga un reporte consolidado con todos los préstamos, fechas de entrega, tiempos límite y retrasos.")

    if st.session_state.prestamos_db:
        excel_bytes = exportar_historial_excel(st.session_state.prestamos_db)
        nombre_reporte = f"Reporte_Prestamos_Laboratorio_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

        st.download_button(
            label="📥 Descargar Reporte Completo en Excel (.xlsx)",
            data=excel_bytes,
            file_name=nombre_reporte,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )
    else:
        st.info("No hay préstamos para exportar.")

# =====================================================================
# PESTAÑA 3: INVENTARIO DE EQUIPOS
# =====================================================================
with tab_inventario:
    st.markdown("### 📊 Inventario de Equipos del Laboratorio")
    st.markdown("Consulta en tiempo real la disponibilidad de los equipos y actualiza la base de datos subiendo un nuevo archivo Excel.")

    inv_col1, inv_col2, inv_col3, inv_col4 = st.columns(4)
    total_eq = len(df_equipos)
    total_cats = df_equipos["categoria"].nunique()
    en_prestamo_total = len(equipos_ocupados)
    disponibles_total = max(total_eq - en_prestamo_total, 0)

    with inv_col1:
        st.metric("Total Equipos Registrados", total_eq)
    with inv_col2:
        st.metric("Categorías / Armarios", total_cats)
    with inv_col3:
        st.metric("Equipos Disponibles", disponibles_total)
    with inv_col4:
        st.metric("Equipos en Préstamo", en_prestamo_total)

    st.markdown("---")

    col_tabla, col_subir = st.columns([3, 2], gap="large")

    with col_tabla:
        st.markdown("##### 📋 Listado Completo")
        
        t_c1, t_c2 = st.columns(2)
        with t_c1:
            filtro_cat_inv = st.selectbox("Filtrar por Armario / Categoría:", ["Todas"] + sorted(df_equipos["categoria"].dropna().unique().tolist()))
        with t_c2:
            buscar_inv = st.text_input("Buscar equipo:", placeholder="Ej: multimetro o resistencia")

        df_mostrar = df_equipos.copy()
        
        estados_dinamicos = []
        for _, r in df_mostrar.iterrows():
            cod = r.get("codigo_equipo", "")
            nom = r.get("nombre_equipo", "")
            etiqueta = f"{cod} — {nom}" if cod else nom
            if etiqueta in equipos_ocupados:
                estados_dinamicos.append("🔴 En Préstamo")
            else:
                estados_dinamicos.append("🟢 Disponible")
        
        df_mostrar["disponibilidad"] = estados_dinamicos

        if filtro_cat_inv != "Todas":
            df_mostrar = df_mostrar[df_mostrar["categoria"] == filtro_cat_inv]

        if buscar_inv.strip():
            q = buscar_inv.strip().lower()
            df_mostrar = df_mostrar[
                df_mostrar["nombre_equipo"].astype(str).str.lower().str.contains(q) |
                df_mostrar["codigo_equipo"].astype(str).str.lower().str.contains(q)
            ]

        st.dataframe(
            df_mostrar[["codigo_equipo", "nombre_equipo", "categoria", "disponibilidad", "ubicacion"]],
            column_config={
                "codigo_equipo": "Código",
                "nombre_equipo": "Descripción del Equipo / Artículo",
                "categoria": "Categoría / Armario",
                "disponibilidad": "Disponibilidad Actual",
                "ubicacion": "Ubicación"
            },
            use_container_width=True,
            hide_index=True
        )

    with col_subir:
        st.markdown("##### 📁 Actualizar Base de Datos")
        st.markdown("""
        Puedes subir un nuevo archivo **Excel (.xlsx)** para actualizar los equipos disponibles.
        El sistema detecta automáticamente la hoja `ArmarioLCD` o formatos con columnas:
        `codigo_equipo`, `nombre_equipo`, `categoria`, `ubicacion`.
        """)

        archivo_subido = st.file_uploader("Seleccionar archivo Excel (.xlsx):", type=["xlsx"])
        if archivo_subido is not None:
            if st.button("💾 Guardar y Aplicar Nuevo Inventario", use_container_width=True):
                with open(NOMBRE_ARCHIVO_EXCEL, "wb") as f:
                    f.write(archivo_subido.getbuffer())
                st.cache_data.clear()
                st.success("✅ ¡Base de datos de inventario actualizada con éxito!")
                st.rerun()

        st.markdown("---")
        st.markdown("##### 📥 Plantilla de Ejemplo")
        st.markdown("Si deseas crear un inventario nuevo desde cero, puedes descargar esta plantilla estándar:")
        
        df_plantilla = pd.DataFrame({
            "codigo_equipo": ["LAB-001", "LAB-002", "LAB-003"],
            "nombre_equipo": ["Osciloscopio Digital", "Multímetro de Banco", "Fuente de Poder"],
            "categoria": ["Medición", "Medición", "Alimentación"],
            "estado": ["Disponible", "Disponible", "Disponible"],
            "ubicacion": ["Estante A1", "Estante B1", "Mesa 1"]
        })
        buffer_plantilla = io.BytesIO()
        df_plantilla.to_excel(buffer_plantilla, index=False)
        
        st.download_button(
            label="Descargar Plantilla Estándar (.xlsx)",
            data=buffer_plantilla.getvalue(),
            file_name="plantilla_equipos_laboratorio.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )

# =====================================================================
# PESTAÑA 4: MENÚ DE CONFIGURACIÓN (TIEMPO, CANTIDADES Y GMAIL)
# =====================================================================
with tab_config:
    st.markdown("### ⚙️ Menú de Configuración de Parámetros y Correo")
    st.markdown("Personaliza el **tiempo de préstamo**, la **cantidad máxima de equipos a prestar** y la cuenta de **Gmail** emisora.")

    col_menu_tiempo, col_menu_correo = st.columns([1, 1], gap="large")

    # -------------------------------------------------------------
    # SUBMENÚ 1: PARÁMETROS DE TIEMPO Y CANTIDAD DE EQUIPOS
    # -------------------------------------------------------------
    with col_menu_tiempo:
        st.markdown("""
        <div class="info-box-dark">
            <h4>⏱️ Parámetros de Préstamo y Cantidad</h4>
            <p style="font-size: 13px; color: #94a3b8;">
                Modifica aquí la duración estándar del préstamo y el número máximo de equipos que se pueden solicitar o prestar.
            </p>
        </div>
        """, unsafe_allow_html=True)

        with st.form("form_config_tiempo_y_cantidades"):
            st.markdown("##### ⏳ Duración y Cantidades")
            
            nueva_duracion = st.number_input(
                "Tiempo estándar del préstamo (Horas):",
                min_value=0.25,
                max_value=24.0,
                value=float(DURACION_PRESTAMO_HORAS),
                step=0.25,
                help="Ejemplo: 1.0 (1 hora), 2.0 (2 horas), 3.5 (3 horas y media), etc."
            )

            nuevo_max_equipos = st.number_input(
                "Cantidad máxima de equipos a prestar por solicitud:",
                min_value=1,
                max_value=50,
                value=int(MAX_EQUIPOS_PERMITIDOS),
                step=1,
                help="Límite máximo de artículos que un estudiante puede seleccionar en su solicitud."
            )

            nuevo_nombre_lab = st.text_input(
                "Nombre del laboratorio / remitente:",
                value=cfg_sistema.get("nombre_laboratorio", "Laboratorio Universitario - Préstamos")
            )

            btn_guardar_parametros = st.form_submit_button("💾 Guardar Parámetros de Préstamo", use_container_width=True)

            if btn_guardar_parametros:
                nueva_cfg_sis = {
                    "duracion_prestamo_horas": float(nueva_duracion),
                    "max_equipos_permitidos": int(nuevo_max_equipos),
                    "nombre_laboratorio": nuevo_nombre_lab.strip()
                }
                guardar_config_sistema(nueva_cfg_sis)
                st.success(f"✅ ¡Guardado! Tiempo: {nueva_duracion}h · Cantidad máxima: {nuevo_max_equipos} equipos.")
                st.rerun()

        st.markdown(f"""
        <div style="background: rgba(56, 189, 248, 0.1); border: 1px solid #0284c7; border-radius: 8px; padding: 14px; margin-top: 16px; font-size: 13px; color: #bae6fd;">
            💡 <strong>Configuración activa en este momento:</strong><br>
            • <strong>Tiempo estándar:</strong> {DURACION_PRESTAMO_HORAS} horas ({int(DURACION_PRESTAMO_HORAS * 60)} min).<br>
            • <strong>Cantidad máxima a prestar:</strong> {MAX_EQUIPOS_PERMITIDOS} equipos por solicitud.<br>
            • <em>Nota:</em> En la pestaña <strong>Control del Encargado</strong> puedes añadir, retirar o modificar la lista de equipos y las horas para cualquier estudiante antes de entregarle el material.
        </div>
        """, unsafe_allow_html=True)

    # -------------------------------------------------------------
    # SUBMENÚ 2: CONFIGURACIÓN DE CORREO GMAIL
    # -------------------------------------------------------------
    with col_menu_correo:
        st.markdown("""
        <div class="info-box-dark">
            <h4>📧 Configuración de Correo (Gmail)</h4>
            <p style="font-size: 13px; color: #94a3b8;">
                Configura la cuenta de Gmail emisora de los comprobantes electrónicos para los estudiantes.
            </p>
        </div>
        """, unsafe_allow_html=True)

        cfg_actual = cargar_config_correo()

        with st.form("form_config_correo_gmail"):
            remitente_in = st.text_input(
                "Cuenta de Gmail Remitente:",
                value=cfg_actual.get("remitente", ""),
                placeholder="tu_laboratorio@gmail.com"
            )
            password_in = st.text_input(
                "Contraseña de Aplicación de Google (16 letras):",
                value=cfg_actual.get("password", ""),
                type="password",
                placeholder="xxxx xxxx xxxx xxxx",
                help="Genera tu clave de 16 caracteres en https://myaccount.google.com/apppasswords"
            )
            nombre_rem_in = st.text_input(
                "Nombre visible del Remitente:",
                value=cfg_actual.get("nombre_remitente", "Laboratorio Universitario")
            )

            btn_guardar_correo = st.form_submit_button("💾 Guardar Configuración de Gmail", use_container_width=True)

            if btn_guardar_correo:
                nueva_cfg_correo = {
                    "proveedor": "gmail",
                    "smtp_server": "smtp.gmail.com",
                    "smtp_port": 465,
                    "usar_ssl": True,
                    "remitente": remitente_in.strip(),
                    "password": password_in.strip().replace(" ", ""),
                    "nombre_remitente": nombre_rem_in.strip()
                }
                guardar_config_correo(nueva_cfg_correo)
                st.success("✅ Configuración de Gmail guardada con éxito.")

        st.markdown("""
        <div style="background: rgba(148, 163, 184, 0.1); border: 1px solid #334155; border-radius: 8px; padding: 14px; margin-top: 14px; font-size: 12px; color: #cbd5e1;">
            <strong>Pasos para la clave de Gmail:</strong><br>
            1. Activa verificación en 2 pasos en Google: <a href="https://myaccount.google.com/security" target="_blank" style="color: #38bdf8;">myaccount.google.com/security</a>.<br>
            2. Genera la clave de 16 letras en: <a href="https://myaccount.google.com/apppasswords" target="_blank" style="color: #38bdf8;">myaccount.google.com/apppasswords</a>.<br>
            3. Pégala arriba sin espacios.
        </div>
        """, unsafe_allow_html=True)

        st.markdown("##### 🧪 Probar Envío con Gmail")
        dest_prueba = st.text_input("Correo para prueba de recepción:", placeholder="tu_correo@ejemplo.com", key="input_test_email")
        
        if st.button("🚀 Enviar Correo de Prueba", use_container_width=True):
            if not dest_prueba.strip():
                st.warning("Ingresa un correo electrónico para la prueba.")
            else:
                with st.spinner("Conectando con Gmail (smtp.gmail.com:465)..."):
                    ok, msg = enviar_correo_confirmacion(
                        destinatario=dest_prueba.strip(),
                        nombre_estudiante="Usuario de Prueba",
                        cedula="1234567890",
                        equipos=["LAB-OSC-01 — Osciloscopio Digital", "LAB-MUL-01 — Multímetro Digital"],
                        id_solicitud=999,
                        duracion_horas=DURACION_PRESTAMO_HORAS
                    )
                    if ok:
                        st.success(f"🎉 ¡Correo de prueba enviado con éxito a **{dest_prueba}**!")
                    else:
                        st.error(f"❌ Error al enviar correo de prueba:\n\n{msg}")